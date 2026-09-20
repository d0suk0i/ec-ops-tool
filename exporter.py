from datetime import datetime
from pathlib import Path
from copy import copy

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ---------- COLOURS ----------

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78"
)

PROFITABLE_FILL = PatternFill(
    "solid",
    fgColor="C6EFCE"
)

HIGH_ROI_FILL = PatternFill(
    "solid",
    fgColor="A9D18E"
)

WARNING_FILL = PatternFill(
    "solid",
    fgColor="FFF2CC"
)

REVIEW_FILL = PatternFill(
    "solid",
    fgColor="FCE4D6"
)

LOSS_FILL = PatternFill(
    "solid",
    fgColor="FFC7CE"
)

ROW_FILL_ODD = PatternFill(
    fill_type="solid",
    start_color="FFFFFF",
    end_color="FFFFFF",
)

ROW_FILL_EVEN = PatternFill(
    fill_type="solid",
    start_color="DCEAF7",
    end_color="DCEAF7",
)

JAPANESE_SHEET_NAMES = {
    "Summary": "サマリー",
    "Product Analysis": "商品分析",
    "Scoring Detail": "スコア詳細",
    "Validation Issues": "入力エラー",
    "Duplicate Issues": "重複データ",
}


JAPANESE_HEADERS = {
    "Metric": "項目",
    "Value": "件数",
    "Rank": "順位",
    "SKU": "SKU",
    "Product": "商品名",
    "Platform": "販売チャネル",
    "Category": "カテゴリ",
    "Quantity": "数量",
    "Sale Price": "販売価格",
    "Item Cost": "商品原価",
    "Shipping": "送料",
    "Other Costs": "その他費用",
    "Platform Fee": "販売手数料",
    "Total Cost": "総コスト",
    "Profit": "利益",
    "Margin": "利益率",
    "ROI": "ROI",
    "Break-Even Price": "損益分岐価格",
    "Status": "ステータス",
    "Profit Score": "利益スコア",
    "Margin Score": "利益率スコア",
    "ROI Score": "ROIスコア",
    "Priority Score": "優先度スコア",
    "Row": "行番号",
    "Issue": "内容",
}

JAPANESE_MINIMUM_WIDTHS = {
    "順位": 9,
    "SKU": 12,
    "商品名": 26,
    "販売チャネル": 20,
    "カテゴリ": 24,
    "数量": 10,
    "販売価格": 14,
    "商品原価": 14,
    "送料": 12,
    "その他費用": 16,
    "販売手数料": 17,
    "総コスト": 14,
    "利益": 14,
    "利益率": 11,
    "ROI": 11,
    "損益分岐価格": 18,
    "ステータス": 42,
    "利益スコア": 14,
    "利益率スコア": 15,
    "ROIスコア": 14,
    "優先度スコア": 18,
    "行番号": 13,
    "内容": 38,
    "項目": 30,
    "件数": 12,
}

JAPANESE_SUMMARY_LABELS = {
    "Products analyzed": "商品数",
    "Profitable products": "黒字商品",
    "Products requiring review": "要確認商品",
    "Loss-making products": "赤字商品",
    "Validation issues": "入力エラー",
    "Duplicate issues": "重複データ",
}


JAPANESE_STATUS_LABELS = {
    "PROFITABLE": "黒字",
    "HIGH ROI": "高ROI",
    "LOW MARGIN": "低利益率",
    "LOSS": "赤字",
    "REVIEW REQUIRED": "要確認",
    "MISSING DATA": "データ不足",
    "INVALID DATA": "無効データ",
    "CONFLICTING DUPLICATE": "SKU重複・内容不一致",
    "UNKNOWN PLATFORM": "未登録販売チャネル",
    "PLATFORM DISABLED": "販売チャネル無効",
    "UNKNOWN CATEGORY": "未登録カテゴリ",
}

# ---------- HELPERS ----------

def format_banded_rows(
    worksheet,
    start_row=2,
    exclude_columns=None,
):
    if exclude_columns is None:
        exclude_columns = set()

    for row_number in range(
        start_row,
        worksheet.max_row + 1,
    ):
        row_fill = (
            ROW_FILL_ODD
            if row_number % 2 == 1
            else ROW_FILL_EVEN
        )

        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            if column_number in exclude_columns:
                continue

            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            # Preserve cells that already have
            # special warning / status fills.
            if (
                cell.fill is not None
                and cell.fill.fill_type == "solid"
                and cell.fill.fgColor.rgb
                not in {
                    "00000000",
                    "000000",
                    "FFFFFFFF",
                    "00FFFFFF",
                }
            ):
                continue

            cell.fill = row_fill

def auto_size_columns(worksheet):
    minimum_widths = {
        "Rank": 10,
        "SKU": 12,
        "Product": 26,
        "Platform": 22,
        "Category": 27,
        "Quantity": 13,
        "Sale Price": 14,
        "Item Cost": 14,
        "Shipping": 13,
        "Other Costs": 13,
        "Platform Fee": 15,
        "Total Cost": 14,
        "Profit": 14,
        "Margin": 11,
        "ROI": 11,
        "Break-Even Price": 17,
        "Status": 42,
        "Row": 10,
        "Issue": 38,
    }

    maximum_width = 45

    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        header = str(column_cells[0].value or "")

        max_length = len(header)

        for cell in column_cells:
            if cell.value is None:
                continue

            if isinstance(cell.value, (int, float)):
                if "¥" in cell.number_format:
                    display_value = f"¥{cell.value:,.2f}"

                elif "%" in cell.number_format:
                    display_value = f"{cell.value:.2f}%"

                else:
                    display_value = str(cell.value)

            else:
                display_value = str(cell.value)

            max_length = max(
                max_length,
                len(display_value),
            )

        calculated_width = max_length + 4

        minimum_width = minimum_widths.get(
            header,
            10,
        )

        final_width = max(
            calculated_width,
            minimum_width,
        )

        final_width = min(
            final_width,
            maximum_width,
        )

        worksheet.column_dimensions[
            column_letter
        ].width = final_width

def format_headers(worksheet):

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D9E1F2"
        )
    )

    for cell in worksheet[1]:

        cell.fill = HEADER_FILL

        cell.font = Font(
            color="FFFFFF",
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.border = thin_border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

def format_status_cells(worksheet, status_column):

    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        cell = worksheet.cell(
            row=row_number,
            column=status_column
        )

        status = str(cell.value or "")

        # Severity takes priority
        if "LOSS" in status:
            cell.fill = LOSS_FILL

        elif "REVIEW REQUIRED" in status:
            cell.fill = REVIEW_FILL

        elif "LOW MARGIN" in status:
            cell.fill = WARNING_FILL

        elif "HIGH ROI" in status:
            cell.fill = HIGH_ROI_FILL

        elif "PROFITABLE" in status:
            cell.fill = PROFITABLE_FILL

def set_data_row_height(
    worksheet,
    height=20,
):
    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        worksheet.row_dimensions[
            row_number
        ].height = height

def format_summary_sheet(worksheet):
    worksheet.auto_filter.ref = None
    worksheet.freeze_panes = None

    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 14

    worksheet.row_dimensions[1].height = 24

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        worksheet.row_dimensions[
            row_number
        ].height = 22

        metric_cell = worksheet.cell(
            row=row_number,
            column=1,
        )

        value_cell = worksheet.cell(
            row=row_number,
            column=2,
        )

        metric = str(
            metric_cell.value or ""
        )

        metric_cell.font = Font(
            bold=True
        )

        value_cell.font = Font(
            bold=True,
            size=12,
        )

        value_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        if metric == "Profitable products":
            row_fill = PROFITABLE_FILL

        elif metric == "Loss-making products":
            row_fill = LOSS_FILL

        elif metric in {
            "Products requiring review",
            "Validation issues",
            "Duplicate issues",
        }:
            row_fill = REVIEW_FILL

        else:
            row_fill = ROW_FILL_EVEN

        metric_cell.fill = row_fill
        value_cell.fill = row_fill

        def format_issue_sheet(
                worksheet,
        ):
            for row_number in range(
                    2,
                    worksheet.max_row + 1,
            ):
                worksheet.row_dimensions[
                    row_number
                ].height = 20

                for column_number in range(
                        1,
                        worksheet.max_column + 1,
                ):
                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                    ).fill = REVIEW_FILL

def format_issue_sheet(
    worksheet,
):
    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        worksheet.row_dimensions[
            row_number
        ].height = 20

        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
            ).fill = REVIEW_FILL

def format_duplicate_issue_sheet(
    worksheet,
):
    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        worksheet.row_dimensions[
            row_number
        ].height = 20

        issue = str(
            worksheet.cell(
                row=row_number,
                column=3,
            ).value or ""
        ).lower()

        if "exact duplicate row removed" in issue:
            row_fill = ROW_FILL_EVEN

        else:
            row_fill = REVIEW_FILL

        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
            ).fill = row_fill

def localize_workbook_to_japanese(
    workbook,
):
    # Translate headers and cell contents
    for worksheet in workbook.worksheets:

        for cell in worksheet[1]:
            if cell.value in JAPANESE_HEADERS:
                cell.value = JAPANESE_HEADERS[
                    cell.value
                ]

        if worksheet.title == "Summary":
            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                cell = worksheet.cell(
                    row=row_number,
                    column=1,
                )

                if cell.value in JAPANESE_SUMMARY_LABELS:
                    cell.value = JAPANESE_SUMMARY_LABELS[
                        cell.value
                    ]

        if worksheet.title == "Product Analysis":
            status_column = None

            for cell in worksheet[1]:
                if cell.value == "ステータス":
                    status_column = cell.column
                    break

            if status_column is not None:
                for row_number in range(
                    2,
                    worksheet.max_row + 1,
                ):
                    cell = worksheet.cell(
                        row=row_number,
                        column=status_column,
                    )

                    if not cell.value:
                        continue

                    parts = str(
                        cell.value
                    ).split(" | ")

                    translated_parts = [
                        JAPANESE_STATUS_LABELS.get(
                            part,
                            part,
                        )
                        for part in parts
                    ]

                    cell.value = " | ".join(
                        translated_parts
                    )

        if worksheet.title == "Validation Issues":
            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                issue_cell = worksheet.cell(
                    row=row_number,
                    column=3,
                )

                issue = str(
                    issue_cell.value or ""
                )

                if issue.startswith(
                    "Missing or invalid "
                ):
                    field_name = issue.replace(
                        "Missing or invalid ",
                        "",
                    )

                    field_names = {
                        "sale_price": "販売価格",
                        "item_cost": "商品原価",
                        "shipping_cost": "送料",
                        "other_costs": "その他費用",
                        "quantity": "数量",
                        "platform": "販売チャネル",
                    }

                    display_name = field_names.get(
                        field_name,
                        field_name,
                    )

                    issue_cell.value = (
                        f"{display_name}が"
                        "未入力または無効"
                    )

        if worksheet.title == "Duplicate Issues":
            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                issue_cell = worksheet.cell(
                    row=row_number,
                    column=3,
                )

                issue = str(
                    issue_cell.value or ""
                )

                if issue == "Exact duplicate row removed":
                    issue_cell.value = (
                        "完全重複行を削除"
                    )

                elif (
                    issue
                    == "Duplicate SKU with conflicting data"
                ):
                    issue_cell.value = (
                        "同一SKUで内容が不一致"
                    )

    # Rename sheet tabs last
    for worksheet in workbook.worksheets:
        if worksheet.title in JAPANESE_SHEET_NAMES:
            worksheet.title = JAPANESE_SHEET_NAMES[
                worksheet.title
            ]

def format_japanese_workbook(
    workbook,
):
    for worksheet in workbook.worksheets:

        # Apply a Japanese-friendly font while
        # preserving existing sizes and emphasis.
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue

                new_font = copy(
                    cell.font
                )

                new_font.name = "Yu Gothic"

                cell.font = new_font

        # Japanese headers look cleaner without
        # the heavy bold weight.
        for cell in worksheet[1]:
            new_font = copy(
                cell.font
            )

            new_font.name = "Yu Gothic"
            new_font.bold = False

            cell.font = new_font

        if worksheet.title == "サマリー":
            for row_number in range(
                    2,
                    worksheet.max_row + 1,
            ):
                for column_number in range(
                        1,
                        worksheet.max_column + 1,
                ):
                    cell = worksheet.cell(
                        row=row_number,
                        column=column_number,
                    )

                    new_font = copy(
                        cell.font
                    )

                    new_font.name = "Yu Gothic"
                    new_font.bold = False

                    cell.font = new_font

        # Recalculate widths after translation.
        for column_cells in worksheet.columns:
            column_letter = (
                column_cells[0].column_letter
            )

            header = str(
                column_cells[0].value or ""
            )

            minimum_width = (
                JAPANESE_MINIMUM_WIDTHS.get(
                    header,
                    10,
                )
            )

            max_length = len(
                header
            ) * 2

            for cell in column_cells:
                if cell.value is None:
                    continue

                text = str(
                    cell.value
                )

                # Japanese/full-width characters need
                # roughly twice the horizontal space.
                display_length = sum(
                    2 if ord(character) > 255
                    else 1
                    for character in text
                )

                max_length = max(
                    max_length,
                    display_length,
                )

            calculated_width = (
                max_length + 3
            )

            final_width = min(
                max(
                    calculated_width,
                    minimum_width,
                ),
                45,
            )

            worksheet.column_dimensions[
                column_letter
            ].width = final_width

# ---------- EXPORT ----------

def export_analysis_report(
    analysis,
    validation_issues,
    duplicate_issues,
    output_dir,
    language="en",
):

    output_dir = Path(output_dir)

    output_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    output_path = (
        output_dir
        / f"ec_analysis_{timestamp}.xlsx"
    )

    # -------------------------
    # SUMMARY
    # -------------------------

    total_products = len(analysis)

    profitable_products = analysis[
        analysis["flags"].str.contains(
            "PROFITABLE",
            na=False
        )
    ].shape[0]

    review_products = analysis[
        analysis["flags"].str.contains(
            "REVIEW REQUIRED",
            na=False
        )
    ].shape[0]

    loss_products = analysis[
        analysis["flags"].str.contains(
            "LOSS",
            na=False
        )
    ].shape[0]

    summary = pd.DataFrame({
        "Metric": [
            "Products analyzed",
            "Profitable products",
            "Products requiring review",
            "Loss-making products",
            "Validation issues",
            "Duplicate issues",
        ],
        "Value": [
            total_products,
            profitable_products,
            review_products,
            loss_products,
            len(validation_issues),
            len(duplicate_issues),
        ],
    })

    # -------------------------
    # CLEAN MAIN REPORT
    # -------------------------

    report_columns = [
        "priority_rank",
        "sku",
        "product_name",
        "platform",
        "category",
        "quantity",
        "sale_price",
        "item_cost",
        "shipping_cost",
        "other_costs",
        "platform_fee",
        "total_cost",
        "profit",
        "margin_pct",
        "roi_pct",
        "break_even_price",
        "flags",
    ]

    product_report = analysis[
        report_columns
    ].copy()

    product_report = product_report.rename(
        columns={
            "priority_rank": "Rank",
            "sku": "SKU",
            "product_name": "Product",
            "platform": "Platform",
            "category": "Category",
            "quantity": "Quantity",
            "sale_price": "Sale Price",
            "item_cost": "Item Cost",
            "shipping_cost": "Shipping",
            "other_costs": "Other Costs",
            "platform_fee": "Platform Fee",
            "total_cost": "Total Cost",
            "profit": "Profit",
            "margin_pct": "Margin",
            "roi_pct": "ROI",
            "break_even_price": "Break-Even Price",
            "flags": "Status",
        }
    )

    # -------------------------
    # SCORING DETAIL
    # -------------------------

    scoring_columns = [
        "sku",
        "product_name",
        "profit_score",
        "margin_score",
        "roi_score",
        "priority_score",
        "priority_rank",
    ]

    scoring_report = analysis[
        scoring_columns
    ].copy()

    scoring_report = scoring_report.rename(
        columns={
            "sku": "SKU",
            "product_name": "Product",
            "profit_score": "Profit Score",
            "margin_score": "Margin Score",
            "roi_score": "ROI Score",
            "priority_score": "Priority Score",
            "priority_rank": "Rank",
        }
    )

    # -------------------------
    # WRITE WORKBOOK
    # -------------------------

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl"
    ) as writer:

        summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        product_report.to_excel(
            writer,
            sheet_name="Product Analysis",
            index=False
        )

        scoring_report.to_excel(
            writer,
            sheet_name="Scoring Detail",
            index=False
        )

        validation_issues.rename(
            columns={
                "row": "Row",
                "sku": "SKU",
                "issue": "Issue",
            }
        ).to_excel(
            writer,
            sheet_name="Validation Issues",
            index=False,
        )

        duplicate_issues.rename(
            columns={
                "row": "Row",
                "sku": "SKU",
                "issue": "Issue",
            }
        ).to_excel(
            writer,
            sheet_name="Duplicate Issues",
            index=False,
        )

        workbook = writer.book

        # -------------------------
        # GENERAL FORMATTING
        # -------------------------

        for worksheet in workbook.worksheets:
            format_headers(
                worksheet
            )

            if worksheet.title not in {
                "Summary",
                "Validation Issues",
                "Duplicate Issues",
            }:
                format_banded_rows(
                    worksheet
                )

                set_data_row_height(
                    worksheet
                )

            auto_size_columns(
                worksheet
            )

        summary_sheet = workbook[
            "Summary"
        ]

        format_summary_sheet(
            summary_sheet
        )

        validation_sheet = workbook[
            "Validation Issues"
        ]

        format_issue_sheet(
            validation_sheet
        )

        duplicate_sheet = workbook[
            "Duplicate Issues"
        ]

        format_duplicate_issue_sheet(
            duplicate_sheet
        )

        # ------------------------
        # PRODUCT ANALYSIS FORMATS
        # ------------------------

        analysis_sheet = workbook[
            "Product Analysis"
        ]

        # -------------------------
        # PRODUCT ANALYSIS FORMATS
        # -------------------------

        analysis_sheet = workbook[
            "Product Analysis"
        ]

        headers = {
            cell.value: cell.column
            for cell in analysis_sheet[1]
        }

        for column_name in [
            "Rank",
            "Quantity",
        ]:
            column_number = headers[
                column_name
            ]

            for row_number in range(
                    2,
                    analysis_sheet.max_row + 1,
            ):
                analysis_sheet.cell(
                    row=row_number,
                    column=column_number,
                ).alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

        currency_columns = [
            "Sale Price",
            "Item Cost",
            "Shipping",
            "Other Costs",
            "Platform Fee",
            "Total Cost",
            "Profit",
            "Break-Even Price",
        ]

        for column_name in currency_columns:

            column_number = headers[
                column_name
            ]

            for row in range(
                2,
                analysis_sheet.max_row + 1
            ):

                analysis_sheet.cell(
                    row=row,
                    column=column_number
                ).number_format = '¥#,##0.00'

        # Margin and ROI are already stored as
        # numbers such as 33.47, so add the %
        # symbol without multiplying by 100.
        for column_name in [
            "Margin",
            "ROI",
        ]:

            column_number = headers[
                column_name
            ]

            for row in range(
                2,
                analysis_sheet.max_row + 1
            ):

                analysis_sheet.cell(
                    row=row,
                    column=column_number
                ).number_format = '0.00"%"'

        # Status cell background colours
        format_status_cells(
            analysis_sheet,
            headers["Status"]
        )

        # Highlight negative profit directly
        profit_column = headers["Profit"]

        for row in range(
            2,
            analysis_sheet.max_row + 1
        ):

            profit_cell = analysis_sheet.cell(
                row=row,
                column=profit_column
            )

            if (
                isinstance(
                    profit_cell.value,
                    (int, float)
                )
                and profit_cell.value < 0
            ):
                profit_cell.fill = LOSS_FILL

        # -------------------------
        # SCORING SHEET FORMATS
        # -------------------------

        scoring_sheet = workbook[
            "Scoring Detail"
        ]

        scoring_headers = {
            cell.value: cell.column
            for cell in scoring_sheet[1]
        }

        for column_name in [
            "Profit Score",
            "Margin Score",
            "ROI Score",
            "Priority Score",
        ]:

            column_number = scoring_headers[
                column_name
            ]

            for row in range(
                2,
                scoring_sheet.max_row + 1
            ):

                scoring_sheet.cell(
                    row=row,
                    column=column_number
                ).number_format = "0.00"

        for column_name in [
            "Profit Score",
            "Margin Score",
            "ROI Score",
            "Priority Score",
            "Rank",
        ]:
            column_number = scoring_headers[
                column_name
            ]

            for row in range(
                    2,
                    scoring_sheet.max_row + 1,
            ):
                scoring_sheet.cell(
                    row=row,
                    column=column_number,
                ).alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

        # Re-run sizing after header names
        # and formatting are finalized.
        for worksheet in workbook.worksheets:

            auto_size_columns(
                worksheet
            )

        if language == "ja":
            localize_workbook_to_japanese(
                workbook
            )

            format_japanese_workbook(
                workbook
            )

    return output_path